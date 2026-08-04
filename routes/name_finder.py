from flask import Blueprint, request, jsonify, session, send_file
import csv
import io
from db import get_conn, put_conn
from routes.contacts import parse_rows

name_finder_bp = Blueprint('name_finder', __name__)


def phone_suffix(phone):
    digits = ''.join(c for c in str(phone) if c.isdigit())
    return digits[-10:] if len(digits) >= 10 else digits


# Phone-suffix -> most-recently-seen contact name, built from every reply
# this tenant has ever exchanged with a name attached. This is the only
# source of a WhatsApp display name we ever have — Meta only hands us a
# contact's name when that person messages the business first, there's no
# lookup API for arbitrary numbers. Matching on the last 10 digits (not the
# full stored string) so a "+91XXXXXXXXXX" in the replies table still
# matches an uploaded "0XXXXXXXXXX" or "XXXXXXXXXX" for the same number.
def get_known_names(user_id):
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("""
            SELECT DISTINCT ON (suffix) suffix, contact_name FROM (
                SELECT RIGHT(regexp_replace(from_phone, '\\D', '', 'g'), 10) AS suffix,
                       contact_name, received_at
                FROM replies
                WHERE user_id = %s AND contact_name IS NOT NULL AND contact_name != ''
            ) sub
            ORDER BY suffix, received_at DESC
        """, (user_id,))
        rows = cur.fetchall()
        cur.close()
        return {r[0]: r[1] for r in rows}
    finally:
        put_conn(conn)


# POST /api/name-finder/parse
# Upload a CSV/Excel of phone numbers (same column-detection as the Contacts
# import), match each against contacts who've already messaged this tenant
# before, and return the enriched list plus a match count. Numbers that have
# never contacted this account come back with no name — there's no data
# source that could ever supply one for those.
@name_finder_bp.route('/parse', methods=['POST'])
def parse_and_match():
    if 'file' not in request.files:
        return jsonify({'error': 'No file provided'}), 400
    user_id = session.get('user_id')

    file = request.files['file']
    fname = (file.filename or '').lower()
    raw = file.read()

    try:
        if fname.endswith('.csv'):
            text = raw.decode('utf-8-sig')
            rows = list(csv.reader(io.StringIO(text)))
        elif fname.endswith(('.xlsx', '.xls')):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
            ws = wb.active
            rows = [
                [str(cell.value) if cell.value is not None else '' for cell in row]
                for row in ws.iter_rows()
            ]
        else:
            return jsonify({'error': 'Only CSV and Excel (.xlsx) files are supported'}), 400

        uploaded = parse_rows(rows)
        known = get_known_names(user_id)

        results = []
        matched = 0
        for c in uploaded:
            found_name = known.get(phone_suffix(c['phone']))
            is_matched = bool(found_name)
            if is_matched:
                matched += 1
            results.append({
                'phone':   c['phone'],
                'name':    c['name'] or found_name or '',
                'matched': is_matched,
            })

        return jsonify({'success': True, 'results': results, 'total': len(results), 'matched': matched})

    except Exception as e:
        return jsonify({'error': str(e)}), 500


# POST /api/name-finder/export
# Regenerates the enriched list (already computed by /parse, sent back from
# the client rather than re-matched) as a downloadable .xlsx.
@name_finder_bp.route('/export', methods=['POST'])
def export_xlsx():
    body = request.get_json() or {}
    results = body.get('results', [])
    if not results:
        return jsonify({'error': 'No data to export'}), 400

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Names'
    ws.append(['Phone', 'Name', 'Matched'])
    for r in results:
        ws.append([r.get('phone', ''), r.get('name', ''), 'Yes' if r.get('matched') else 'No'])
    for col in ws.columns:
        width = max(len(str(cell.value)) for cell in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(width, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='names.xlsx'
    )
